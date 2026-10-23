"""
Jev와 System One 결정 모델 — 발표자료(PPTX) 및 슬라이드 아웃라인(MD) 생성 스크립트

실행:  uv run --with python-pptx python3 build_pptx.py
산출:  jev-decision-models-presentation.pptx, jev-decision-models-presentation.md
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- 스타일
NAVY = "12263F"; BLUE = "1F4E79"; MID = "4A7FB5"; PALE = "EEF4FB"; AMBER = "E8A33D"
TEAL = "2A9D8F"; RED = "C0634A"; GREY = "5B6770"; LIGHT = "F5F8FC"; WHITE = "FFFFFF"; DARK = "1B1F24"
PURPLE = "8A6FAE"; GREEN = "5A9A5A"
FONT = "Malgun Gothic"; MONO = "Consolas"
FOOT = "Jev와 System One 모델  |  2026-10-23"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDES = []   # md 출력용 메타


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def txt(slide, x, y, w, h, text, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.1):
    """text: 문자열 또는 문자열 리스트(문단)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = text if isinstance(text, list) else [text]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        r = para.add_run()
        r.text = p
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = rgb(color)
    return tb


def box(slide, x, y, w, h, fill=PALE, line=MID, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def card(slide, x, y, w, h, title, body=None, fill=PALE, line=MID, tcolor=BLUE, tsize=16, bsize=13,
         bcolor=DARK):
    box(slide, x, y, w, h, fill, line)
    txt(slide, x + 0.12, y + 0.08, w - 0.24, 0.5, title, tsize, True, tcolor)
    if body:
        txt(slide, x + 0.12, y + 0.58, w - 0.24, h - 0.65, body, bsize, False, bcolor)


def chip(slide, x, y, w, h, text, fill=BLUE, color=WHITE, size=13, bold=True):
    s = box(slide, x, y, w, h, fill, None)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text if isinstance(text, list) else [text]
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, w=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def code(slide, x, y, w, h, lines, size=11):
    box(slide, x, y, w, h, "1E2A3A", None)
    txt(slide, x + 0.1, y + 0.06, w - 0.2, h - 0.1, lines, size, False, "E6EDF5", font=MONO, line_spacing=1.05)


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.4, header_fill=BLUE, first_bold=True, hl=None):
    nrows, ncols = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    t = gs.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(w * cw / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = FONT
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = rgb(header_fill)
                r.font.bold = True
                r.font.color.rgb = rgb(WHITE)
            else:
                c.fill.fore_color.rgb = rgb("DCE8F5" if i % 2 == 1 else WHITE)
                r.font.color.rgb = rgb(DARK)
                if j == 0 and first_bold:
                    r.font.bold = True
                if hl and (i, j) in hl:
                    r.font.bold = True
                    r.font.color.rgb = rgb(hl[(i, j)])
    return t


def base(title, kicker=None, num=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    box(s, 0, 0, 0.18, 7.5, BLUE, None, MSO_SHAPE.RECTANGLE)
    txt(s, 0.55, 0.32, 12.2, 0.8, title, 28, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        txt(s, 0.55, 1.05, 12.2, 0.4, kicker, 14, False, GREY)
    txt(s, 0.55, 7.05, 9, 0.3, FOOT, 10, False, GREY)
    txt(s, 12.0, 7.05, 0.8, 0.3, str(len(prs.slides)), 10, False, GREY, align=PP_ALIGN.RIGHT)
    return s


LAST_NOTE = [""]


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
    LAST_NOTE[0] = text


def meta(title, visual, points, note=None):
    """아웃라인 MD용 메타. 스피커 노트는 직전 notes() 호출의 전문을 사용한다."""
    SLIDES.append(dict(title=title, visual=visual, points=points, notes=LAST_NOTE[0]))


# ================================================================ 슬라이드
# ---- 1. 표지
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
box(s, 0, 0, 0.35, 7.5, AMBER, None, MSO_SHAPE.RECTANGLE)
txt(s, 0.9, 1.5, 11.5, 1.2, "Jev와 System One 모델", 48, True, WHITE)
txt(s, 0.9, 2.65, 11.5, 0.9, "선택·판단 전용 모델의 입출력 사양과 실행 방식", 26, False, "BFD3EA")
chip(s, 0.9, 4.2, 2.4, 0.8, ["state", "상태 + 질문"], MID, WHITE, 16)
arrow(s, 3.35, 4.6, 4.45, 4.6, AMBER, 3)
chip(s, 4.5, 4.2, 2.4, 0.8, ["모델", "prefill 1회"], AMBER, NAVY, 16)
arrow(s, 6.95, 4.6, 8.05, 4.6, AMBER, 3)
chip(s, 8.1, 4.2, 3.2, 0.8, ["확률분포", "choice · score · noul"], TEAL, WHITE, 16)
txt(s, 0.9, 6.3, 11, 0.4, "2026-10-23 기준  |  1차 자료와 집계 자료를 구분해 표기", 14, False, "9FB6D0")
notes(s, "오늘 발표는 2026년 9월 15일 TypeSafe AI가 공개한 Jev와, 이후 5주 동안 쏟아진 선택·판단 전용 모델들을 다룹니다. 초점은 세 가지입니다. 입출력 형식, 내부 실행 방식, 그리고 실제로 돌리는 방법입니다. 모든 수치는 2026년 10월 23일 기준이며, 벤더가 직접 확인한 사실과 집계 사이트의 전언을 구분해서 말씀드리겠습니다.")
meta("Jev와 System One 모델", "남색 배경. 왼쪽부터 state → 모델(prefill 1회) → 확률분포(choice·score·noul)로 이어지는 3단 플로우.",
     ["생성하지 않는 모델", "입출력 · 실행 방식 · 실행 방법", "2026-10-23 기준"],
     "9월 15일 Jev 공개 이후 5주간의 선택·판단 전용 모델을 다룬다. 벤더 확인 사실과 집계 사이트 전언을 구분해 표기한다.")

# ---- 2. 아젠다
s = base("오늘의 흐름")
items = [("01", "입출력 사양", "실제 JSON 계약"), ("02", "실행 방식", "한 번의 forward pass"),
         ("03", "순서 의존성", "Runoff"), ("04", "지형과 성능", "모델 · 벤치마크"),
         ("05", "실전 패턴", "게이팅 · 캐스케이드"), ("06", "모르는 것", "비공개 영역")]
for i, (n, t, d) in enumerate(items):
    cx = 0.7 + (i % 3) * 4.15
    cy = 1.7 + (i // 3) * 2.5
    box(s, cx, cy, 3.85, 2.1, PALE, MID)
    txt(s, cx + 0.2, cy + 0.15, 1.2, 0.8, n, 36, True, AMBER)
    txt(s, cx + 0.2, cy + 0.95, 3.4, 0.5, t, 22, True, NAVY)
    txt(s, cx + 0.2, cy + 1.45, 3.4, 0.5, d, 14, False, GREY)
notes(s, "여섯 개의 묶음으로 진행합니다. 앞의 세 묶음이 이 모델들이 무엇을 받고 무엇을 계산하는지에 대한 이야기이고, 뒤의 세 묶음이 어떤 모델이 있고 어떻게 쓰며 무엇을 아직 모르는지에 대한 이야기입니다.")
meta("오늘의 흐름", "2×3 카드 그리드. 번호(01~06), 제목, 부제.", ["입출력 사양", "실행 방식", "순서 의존성", "지형과 성능", "실전 패턴", "모르는 것"],
     "앞 세 묶음은 모델이 무엇을 받고 계산하는지, 뒤 세 묶음은 어떤 모델이 있고 어떻게 쓰며 무엇을 모르는지를 다룬다.")

# ---- 3. 3줄 요약
s = base("결론 먼저 — 3줄")
rows3 = [("1", "계약은 수렴했다", "state + questions(choice·score·noul) → probabilities·confidence. Jev JSON을 다수 벤더가 호환 구현, OpenAI만 변형"),
         ("2", "단일 패스는 약식이다", "디코딩 없이 prefill 1회 + 선택지 판독. 구현은 어휘 판독 · 결정 헤드 · 대조 점수화 · 제한적 사고. Jev 내부는 비공개"),
         ("3", "순서와 과신이 약점이다", "선택지 순서만 바꿔도 뒤집힘 → Runoff. 지식·추론에서 과신 → 임계값은 자기 데이터로")]
for i, (n, h, d) in enumerate(rows3):
    y = 1.5 + i * 1.8
    box(s, 0.7, y, 11.9, 1.55, PALE, MID)
    txt(s, 0.9, y + 0.15, 1.0, 1.2, n, 54, True, AMBER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 2.1, y + 0.15, 10.2, 0.5, h, 22, True, NAVY)
    txt(s, 2.1, y + 0.72, 10.2, 0.8, d, 14, False, DARK)
notes(s, "결론부터 말씀드립니다. 첫째, 입출력 계약은 사실상 수렴했습니다. 둘째, 한 번의 forward pass라는 표현은 약식이며 내부 구현은 최소 세 가지 계열과 변종으로 나뉩니다. Jev 자체의 내부는 공개되지 않았습니다. 셋째, 선택지 순서 의존성과 과신이 실무의 두 가지 약점입니다.")
meta("결론 먼저 — 3줄", "큰 숫자 1/2/3과 한 문장 요약 카드 세 장.", ["계약은 수렴했다", "단일 패스는 약식이다", "순서와 과신이 약점이다"],
     "입출력 계약 수렴, 단일 패스의 약식 표현, 순서 의존성과 과신이라는 두 약점을 결론으로 먼저 제시한다.")

# ---- 4. System One이란
s = base("System One 모델이란", "텍스트를 쓰지 않고, 정해진 답 위의 확률분포를 낸다")
txt(s, 0.7, 1.6, 6, 0.4, "LLM 구조화 출력 (자기회귀)", 16, True, RED)
xs = [0.7, 2.55, 4.4, 6.25]
labs = ["prefill", "토큰 1", "토큰 … N", "JSON 검증"]
for i, (x, l) in enumerate(zip(xs, labs)):
    chip(s, x, 2.1, 1.55, 0.75, l, "F3D5CC", RED, 14)
    if i < 3:
        arrow(s, x + 1.57, 2.48, xs[i + 1] - 0.02, 2.48, RED, 2)
txt(s, 8.3, 2.0, 4.6, 1.0, ["지연 ∝ 출력 길이", "확률은 logprobs를 따로 요청", "스키마 오류 가능"], 13, False, GREY)
txt(s, 0.7, 3.5, 6, 0.4, "System One (단일 패스)", 16, True, BLUE)
xs2 = [0.7, 3.2, 5.7, 8.2]
labs2 = ["prefill 1회\nstate+질문+선택지", "선택지 위치\n점수 판독", "softmax\n확률분포", "타입 보장된 답"]
for i, (x, l) in enumerate(zip(xs2, labs2)):
    chip(s, x, 4.0, 2.2, 0.95, l.split("\n"), PALE, BLUE, 14)
    if i < 3:
        arrow(s, x + 2.22, 4.48, xs2[i + 1] - 0.02, 4.48, BLUE, 2)
for i, (h, d) in enumerate([("출력이 공짜", "비용은 전부 입력(prefill)"), ("스키마 오류 없음", "선택지 집합이 사전 고정"), ("모든 후보에 확률", "임계값·기권 가능")]):
    card(s, 0.7 + i * 4.1, 5.35, 3.9, 1.4, h, d, PALE, MID, BLUE, 16, 13)
notes(s, "LLM이 구조화된 답을 내려면 prefill 이후 토큰을 하나씩 샘플링합니다. 지연은 출력 길이에 비례하고, 스키마 오류가 날 수 있습니다. System One 모델은 prefill 한 번 뒤 선택지 위치의 점수를 읽어 softmax를 취합니다. 그래서 출력이 사실상 무료이고, 모든 후보에 확률이 붙습니다. 다만 환각할 수 없다는 말은 스키마상 존재하지 않는 값을 낼 수 없다는 뜻이지 고른 답이 사실이라는 보장이 아닙니다.")
meta("System One 모델이란", "위: LLM 자기회귀 체인(prefill→토큰들→JSON 검증, 붉은색). 아래: 단일 패스 체인(prefill 1회→점수 판독→softmax→타입 보장, 파란색). 하단 3개 카드.",
     ["출력이 공짜", "스키마 오류 없음", "모든 후보에 확률"],
     "디코딩이 없어 비용은 입력에 있고 스키마 오류가 구조적으로 없다. 환각할 수 없다는 말은 값의 형식에 대한 것이지 사실성에 대한 것이 아니다.")

# ---- 5. 타임라인
s = base("5주 만에 30여 개", "2026-09-15 ~ 10-22 · 날짜 일부는 집계 사이트 표기")
box(s, 0.7, 3.85, 11.9, 0.08, GREY, None, MSO_SHAPE.RECTANGLE)
ev = [("09-15", "Jev 1.13", AMBER, 0), ("09-24", "GLiNER2.5\nDecide", MID, 1), ("09-30", "GLiDE\nd1:free", MID, 0),
      ("10-01", "Clef\npplx-decider", MID, 1), ("10-06", "OpenAI\nDecisions", RED, 0), ("10-07", "d1-3B\nd1-omni", MID, 1),
      ("10-13", "AA Index\nRunoff 논문", TEAL, 0), ("10-14", "Jev 1.14", AMBER, 1), ("10-19", "Clairic\nCS1 논문", TEAL, 0), ("10-22", "GLiDE 1.2", MID, 1)]
for i, (d, l, c, up) in enumerate(ev):
    x = 0.8 + i * 1.2
    dot = box(s, x + 0.35, 3.78, 0.22, 0.22, c, None, MSO_SHAPE.OVAL)
    yy = 2.6 if up == 0 else 4.3
    txt(s, x - 0.1, yy, 1.1, 0.35, d, 13, True, c, align=PP_ALIGN.CENTER)
    txt(s, x - 0.15, yy + 0.35, 1.2, 0.8, l.split("\n"), 12, False, DARK, align=PP_ALIGN.CENTER)
txt(s, 0.7, 5.8, 11.9, 0.9, ["공통 배경: Jev가 JSON 계약을 공개 → 호환 구현의 진입 장벽 하락", "Qwen3.5/3.8 · Gemma 4 위에 LoRA + 작은 헤드만으로도 구현 가능"], 15, False, NAVY)
notes(s, "타임라인입니다. 날짜는 벤더 공지가 확인된 것과 집계 사이트 표기에 의존한 것이 섞여 있습니다. 5주 동안 호스팅 모델과 오픈 가중치 모델이 동시에 쏟아졌습니다. 배경은 두 가지입니다. 하나는 Jev가 JSON 계약을 공개해 호환 구현이 쉬워졌다는 점, 다른 하나는 오픈 백본 위에 LoRA와 작은 헤드만 얹어도 만들 수 있다는 점입니다.")
meta("5주 만에 30여 개", "가로 타임라인. 위/아래 교차로 10개 이벤트 점(Jev 1.13, GLiNER2.5-Decide, GLiDE·d1, Clef·pplx, OpenAI, d1-3B, AA·Runoff, Jev 1.14, Clairic·CS1, GLiDE 1.2).",
     ["09-15 Jev 공개", "10-06 OpenAI 공개 베타", "10-13 AA 지수 · Runoff 논문"],
     "날짜는 벤더 공지 확인분과 집계 사이트 표기가 섞여 있다. 호환 JSON 계약과 오픈 백본 + LoRA 구조가 확산의 배경이다.")

# ---- 6. DIV
def divider(num, title, sub):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    box(s, 0, 0, 0.35, 7.5, AMBER, None, MSO_SHAPE.RECTANGLE)
    txt(s, 1.0, 2.0, 3, 1.6, num, 96, True, AMBER)
    txt(s, 1.0, 3.9, 11, 1, title, 40, True, WHITE)
    txt(s, 1.0, 4.9, 11, 0.6, sub, 20, False, "BFD3EA")
    return s

s = divider("01", "입출력 사양", "실제 JSON으로 보는 벤더별 계약")
notes(s, "첫 번째 묶음은 입출력 사양입니다.")
meta("01 입출력 사양", "구간 표지(남색 배경, 큰 번호).", ["실제 JSON으로 보는 벤더별 계약"], "입출력 사양 구간의 시작.")

# ---- 7. 공통 계약
s = base("공통 계약 — state + 이름 붙은 질문", "질문은 서로 독립 평가 · 한 요청에 여러 개")
chip(s, 0.7, 1.7, 2.6, 0.8, ["state", "문자열 · JSON · 배열"], MID, WHITE, 15)
arrow(s, 3.35, 2.1, 4.1, 2.1, GREY, 2.5)
chip(s, 4.15, 1.7, 3.0, 0.8, ["questions", "이름 → {type, instructions, criteria}"], BLUE, WHITE, 14)
arrow(s, 7.2, 2.1, 7.95, 2.1, GREY, 2.5)
chip(s, 8.0, 1.7, 4.6, 0.8, ["answers", "이름별 probabilities · confidence"], TEAL, WHITE, 15)
prim = [("noul", "참 / 거짓", "criteria: true·false 설명(선택)", "P(true) 하나\n별도 confidence 없음"),
        ("choice", "후보 중 하나", "criteria: 라벨 → 설명 맵\n최대 255개", "choice · probabilities\nconfidence"),
        ("score", "순서 있는 척도", "criteria: 레벨 설명 배열\n최대 10단계", "score(기댓값) · legend\nprobabilities · confidence")]
for i, (n, q, inn, out) in enumerate(prim):
    x = 0.7 + i * 4.1
    box(s, x, 3.0, 3.9, 3.6, PALE, MID)
    txt(s, x + 0.2, 3.1, 3.5, 0.5, n, 24, True, BLUE)
    txt(s, x + 0.2, 3.65, 3.5, 0.4, q, 15, True, AMBER)
    txt(s, x + 0.2, 4.2, 3.5, 0.3, "넣는 것", 12, True, GREY)
    txt(s, x + 0.2, 4.5, 3.5, 0.8, inn.split("\n"), 13, False, DARK)
    txt(s, x + 0.2, 5.4, 3.5, 0.3, "돌려받는 것", 12, True, GREY)
    txt(s, x + 0.2, 5.7, 3.5, 0.8, out.split("\n"), 13, False, DARK)
notes(s, "모든 벤더가 공유하는 골격은 같습니다. state 하나에 이름 붙은 질문 여럿을 붙이고, 질문마다 독립적으로 답을 받습니다. primitive는 세 가지입니다. noul은 참거짓, choice는 후보 중 하나, score는 순서 있는 척도입니다. OpenAI는 noul을 predicate라고 부릅니다.")
meta("공통 계약 — state + 이름 붙은 질문", "상단: state → questions → answers 3단 칩. 하단: noul / choice / score 세 카드(넣는 것, 돌려받는 것).",
     ["noul = 참/거짓", "choice = 후보 중 하나 (≤255)", "score = 순서 척도 (≤10)"],
     "질문은 서로의 답을 볼 수 없고 독립 평가된다. OpenAI는 noul을 predicate라 부른다.")

# ---- 8. Jev 요청/응답
s = base("Jev — 요청과 응답", "POST https://api.typesafe.ai/v1/systemone · Bearer 키")
code(s, 0.7, 1.6, 6.0, 5.2, [
    '{ "model": "jev-1.13.0",',
    '  "state": { "ticket": "결제가 두 번 됐어요..." },',
    '  "questions": {',
    '    "team": { "type": "choice",',
    '      "instructions": "처리할 팀은?",',
    '      "criteria": {',
    '        "billing": "결제, 환불",',
    '        "technical": "앱 오류",',
    '        "fraud": "무단 사용" } },',
    '    "urgency": { "type": "score",',
    '      "instructions": "얼마나 급한가?",',
    '      "criteria": ["기다려도 됨",',
    '        "오늘 안에", "지금 막혀 있음"] },',
    '    "refund": { "type": "noul",',
    '      "instructions": "환불 요청인가?" }',
    '  } }'], 11)
code(s, 6.9, 1.6, 5.8, 3.55, [
    '{ "model": "jev-1.13.0",',
    '  "answers": {',
    '    "team": { "type": "choice",',
    '      "choice": "billing", "confidence": 0.895,',
    '      "probabilities": {',
    '        "billing": 0.93, "technical": 0.04,',
    '        "fraud": 0.03 } },',
    '    "urgency": { "type": "score",',
    '      "score": 1.65, "confidence": 0.475,',
    '      "probabilities": {"0":0.05,"1":0.25,"2":0.70} },',
    '    "refund": { "type": "noul", "noul": 0.97 }',
    '  } }'], 11)
card(s, 6.9, 5.3, 5.8, 1.5, "예시는 스키마에 맞춰 직접 구성한 값", "실제 호출 결과 아님\n질문 수 상한은 문서에 명시 없음", LIGHT, MID, BLUE, 14, 12)
notes(s, "Jev의 요청과 응답 예시입니다. 문서의 스키마에 맞춰 직접 구성한 값이며 실제 호출 결과가 아닙니다. 실제 응답에는 usage 필드가 붙습니다. 질문 수의 상한은 문서에 명시되어 있지 않고, choice는 255개, score는 10단계가 상한입니다.")
meta("Jev — 요청과 응답", "좌: 요청 JSON 코드 박스. 우상: 응답 JSON 코드 박스. 우하: '직접 구성한 예시' 주의 카드.",
     ["model + state + questions", "answers.{이름}.probabilities", "예시는 구성한 값"],
     "문서의 스키마에 맞춰 구성한 예시이며 실제 호출 결과가 아니다. 질문 수 상한은 명시되지 않았고 choice 255, score 10단계가 상한이다.")

# ---- 9. 벤더 비교
s = base("벤더별 계약 비교", "Jev 계약이 사실상 표준 · OpenAI는 변형")
rows = [["", "Jev", "OpenAI Decisions", "Liquid d1", "GLiNER2.5-Decide", "Clairic"],
        ["입력 키", "state", "input (+메시지 배열)", "state + images[]", "텍스트 + 라벨", "state + images[]"],
        ["질문 컨테이너", "questions 맵", "questions 배열(name)", "맵", "스키마 작업명", "fields 맵"],
        ["참/거짓", "noul", "predicate", "noul", "yes/no 라벨", "—"],
        ["choice 후보", "criteria 맵 ≤255", "choices[{value,desc}]", "맵", "라벨 리스트", "—"],
        ["score", "배열 ≤10", "levels[{label,desc}]", "배열 2–10", "'0'–'10' 라벨", "estimate"],
        ["기본 출력", "확률 + confidence", "확률 배열 + confidence", "확률 + confidence", "라벨 문자열", "value·p_null"],
        ["호출", "REST", "REST", "REST / 로컬", "로컬 라이브러리", "Workers AI"]]
table(s, 0.6, 1.6, 12.1, rows, [1.5, 1.7, 2.1, 1.7, 1.9, 1.7], 12, 0.58)
txt(s, 0.6, 6.4, 12.1, 0.5, "호환: Cloudflare Clef · Fastino GLiDE · Perplexity · Respan · Upstage · Inception 모두 state/questions 본문 수용", 13, True, TEAL)
notes(s, "다섯 계약을 비교했습니다. Jev, Liquid d1, Clef, GLiDE, Perplexity 등은 같은 state와 questions 본문을 받습니다. OpenAI는 input과 배열형 questions, predicate라는 이름을 씁니다. GLiNER2.5-Decide는 서버 API가 아니라 로컬 라이브러리이며 기본 출력이 라벨 문자열입니다. Clairic은 선택지를 고르는 대신 타입 있는 값을 추출합니다.")
meta("벤더별 계약 비교", "6열 비교 표(Jev, OpenAI, d1, GLiNER2.5-Decide, Clairic). 하단에 호환 벤더 목록.", ["Jev 계약이 사실상 표준", "OpenAI는 배열형 변형", "Clairic은 추출 모델"],
     "벤더를 섞어 쓰려면 OpenAI용 어댑터가 필요하다. 호환 구현은 base_url만 바꿔 교체 가능하다.")

# ---- 10. 숫자 읽는 법
s = base("숫자를 읽는 법", "score는 기댓값, confidence는 분포의 뾰족함")
card(s, 0.7, 1.6, 5.9, 2.5, "score = Σ i · p_i", "레벨 확률 (0.1, 0.7, 0.2)\n→ 0·0.1 + 1·0.7 + 2·0.2 = 1.1\n서로 다른 분포가 같은 점수를 낼 수 있다", PALE, MID, BLUE, 18, 17)
card(s, 6.8, 1.6, 5.9, 2.5, "confidence 수식", "choice: (p_max − 1/n) / (1 − 1/n)\nnoul: |2p − 1| (GLiDE)\nscore: 1 − Σ p_i|i−m| / MAD_unif", PALE, MID, BLUE, 18, 17)
box(s, 0.7, 4.4, 12.0, 2.3, "FDF6E8", AMBER)
txt(s, 0.95, 4.5, 11.5, 0.5, "OpenAI 문서 예시에 대입하면 (이 글의 계산)", 16, True, "7A5710")
txt(s, 0.95, 5.05, 11.5, 1.6, ["choice (0.95, 0.02, 0.01, 0.02) → (0.95−0.25)/0.75 = 0.933  (문서: 0.93)",
                              "score (0.1, 0.7, 0.2) → 1 − 0.3/(2/3) = 0.55  (문서: 0.55)",
                              "→ 두 예시 모두 일치. 벤더 간 confidence 비교가 크게 어긋나지 않는다"], 14, False, DARK)
notes(s, "score는 레벨 번호와 확률의 곱의 합, 즉 기댓값입니다. 서로 다른 분포가 같은 값을 낼 수 있으므로 probabilities와 confidence를 함께 읽어야 합니다. TypeSafe 문서의 confidence 수식을 OpenAI 문서 예시에 대입하면 두 값이 모두 문서의 숫자와 일치합니다. OpenAI가 수식을 공개한 것은 아니므로 같은 정의라고 단정할 수는 없지만, 임계값을 벤더 간에 옮겨도 의미가 크게 어긋나지 않는다는 뜻입니다.")
meta("숫자를 읽는 법", "상단 2카드(score 기댓값, confidence 수식 3종). 하단 호박색 박스에 OpenAI 예시 대입 검산.", ["score = 확률 가중 기댓값", "confidence는 분포의 뾰족함", "OpenAI 예시와 일치"],
     "OpenAI는 수식을 공개하지 않았으므로 같은 정의라 단정할 수는 없지만 예시 두 개가 모두 일치한다.")

# ---- 11. DIV
s = divider("02", "실행 방식", "‘한 번의 forward pass’의 실체")
notes(s, "두 번째 묶음은 실행 방식입니다.")
meta("02 실행 방식", "구간 표지.", ["한 번의 forward pass의 실체"], "실행 방식 구간의 시작.")

# ---- 12. forward pass 실체
s = base("‘한 번의 forward pass’의 실제", "요청 → 응답 데이터 흐름")
flow = ["애플리케이션\nstate+questions", "API 계층\n검증·토큰 계산", "질문별 프롬프트\nstate+지시+선택지", "배치·패킹\n패딩 없이", "prefill 1회\nhidden state", "판독\n어휘·헤드·유사도", "softmax\n+ 온도 보정", "answers\n+ usage"]
for i, l in enumerate(flow):
    x = 0.55 + i * 1.57
    chip(s, x, 1.9, 1.4, 1.0, l.split("\n"), PALE if i not in (4, 5) else AMBER, BLUE if i not in (4, 5) else NAVY, 11)
    if i < 7:
        arrow(s, x + 1.41, 2.4, x + 1.56, 2.4, GREY, 1.8)
tt = [("출력이 공짜인 이유", "디코딩이 없어 비용이 전부 prefill\nAA도 입력 토큰만으로 비용 계산"),
      ("지연이 입력에 좌우", "d1-3B 4090: 짧은 질문 8 ms\n3.4K 토큰 state 102 ms"),
      ("‘단일 패스’의 단위는 제각각", "질문별 패스 (Liquid 호스팅·GLiDE)\n공유·결합 (Clef joint head · 캐시)")]
for i, (h, d) in enumerate(tt):
    card(s, 0.7 + i * 4.1, 3.5, 3.9, 2.6, h, d.split("\n"), PALE, MID, BLUE, 17, 16)
txt(s, 0.7, 6.3, 12, 0.5, "Jev의 usage.output_tokens는 질문당 약 20 — 0이 아니다 (설명 없음)", 15, True, RED)
notes(s, "데이터 흐름입니다. 핵심은 prefill 한 번 뒤 선택지 위치에서 점수를 읽는 것입니다. 출력이 무료인 이유, 지연이 입력 길이에 좌우되는 이유가 여기서 나옵니다. 다만 단일 패스의 단위는 벤더마다 다릅니다. Liquid 호스팅과 GLiDE는 질문마다 별도 패스이고, Clef는 모든 질문의 선택지를 함께 점수화합니다. 그리고 Jev는 출력 토큰이 무료라고 하면서도 usage에 질문당 약 20 정도의 출력 토큰을 보고합니다. 이 불일치는 벤더가 설명한 적이 없습니다.")
meta("‘한 번의 forward pass’의 실제", "8단 플로우(앱→API→프롬프트 구성→패킹→prefill→판독→softmax→응답). 하단 카드 3장과 경고 문구.",
     ["디코딩 없음 → 출력 무료", "지연 ∝ 입력 길이", "패스 단위는 벤더마다 다르다"],
     "Jev의 usage.output_tokens가 질문당 약 20으로 일정하나 의미는 설명되지 않았다. 이 글의 관찰이며 확인되지 않았다.")

# ---- 13. 네 계열
s = base("네 가지 실행 계열", "판독을 어디서 하는가")
fam = [("A. 어휘 판독", "선택지를 문자로 제시\n다음 토큰 확률을 선택지에 대해 재정규화", "Tev1(문자 1개 생성)\n논문의 listwise 정식화", MID),
       ("B. 결정 헤드", "hidden state 위 작은 헤드가\n선택지별 logit 직접 출력", "Clef · Kev · Strands\nVon · GLiNER2.5-Decide · d1-omni", TEAL),
       ("C. 대조 점수화", "상태·후보 행동을 각각 임베딩\n유사도 softmax", "CLM-8B\n(행동 임베딩 캐시)", PURPLE),
       ("D. 제한적 사고", "불확실하면 짧게 생각 후\n재판독·평균", "Span-02 (확률 < 0.8)\nGLiDE 런치판", AMBER)]
for i, (h, d, ex, c) in enumerate(fam):
    x = 0.6 + i * 3.1
    box(s, x, 1.7, 2.95, 4.6, WHITE, c, lw=2.5)
    box(s, x, 1.7, 2.95, 0.75, c, None)
    txt(s, x + 0.1, 1.78, 2.75, 0.6, h, 17, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + 0.15, 2.6, 2.65, 1.7, d.split("\n"), 15, False, DARK)
    txt(s, x + 0.15, 4.4, 2.65, 0.3, "대표", 11, True, GREY)
    txt(s, x + 0.15, 4.7, 2.65, 1.5, ex.split("\n"), 15, True, NAVY)
txt(s, 0.6, 6.45, 12.2, 0.5, "‘출력 토큰 0’은 업계 공통 사실이 아니다 — Tev1은 문자 1개를 생성하고, Span-02·GLiDE 런치판은 사고한다", 14, True, RED)
notes(s, "실행 방식은 네 계열로 나뉩니다. A는 디코더 언어모델의 어휘 판독입니다. B는 백본의 hidden state 위 작은 헤드가 선택지를 직접 점수화하는 방식으로 가장 많은 모델이 여기에 속합니다. C는 대조 학습 기반 점수화입니다. D는 불확실할 때만 짧게 생각하는 변종입니다. 따라서 출력 토큰이 0이라는 서술은 업계 공통 사실이 아닙니다.")
meta("네 가지 실행 계열", "4열 카드(A 어휘 판독·B 결정 헤드·C 대조 점수화·D 제한적 사고). 각 카드에 설명과 대표 모델.", ["A 어휘 판독", "B 결정 헤드", "C 대조 점수화", "D 제한적 사고"],
     "출력 토큰 0이 업계 공통 사실은 아니다. Tev1은 문자 하나를 생성하고 Span-02와 GLiDE 런치판은 사고한다.")

# ---- 14. 모델별 구조
s = base("모델별 구조 — 확인된 만큼만", "벤더/카드 확인 vs 집계 인용")
rows = [["모델", "구조", "판독", "근거"],
        ["Clef 27B", "Qwen3.8-27B + joint schema head", "질문·선택지 결합 점수화", "카드·블로그"],
        ["d1-3B", "LFM2.5-VL-3B 디코더 VLM", "카드에 없음", "카드"],
        ["d1-omni-600M", "양방향 인코더 + 비전·오디오", "결정 헤드 + 유형별 온도", "카드"],
        ["GLiNER2.5-Decide", "DeBERTa-v3-large 340M", "라벨을 호출 시 입력, 다중 헤드", "카드"],
        ["Kev", "Qwen + LoRA r16 + 포인터 헤드", "결정 토큰 대 선택지, 온도 동봉", "집계"],
        ["Von", "ModernBERT 395M", "선택지 마커를 한 번에 점수화", "집계"],
        ["CLM-8B", "동결 Qwen3-8B + 프로젝션", "InfoNCE 유사도 softmax", "집계"],
        ["Jev", "비공개", "비공개", "—"]]
table(s, 0.6, 1.6, 12.1, rows, [1.7, 3.5, 3.6, 1.2], 13, 0.56, hl={(8, 1): RED, (8, 2): RED})
notes(s, "모델별 구조를 정리했습니다. 근거 열에서 카드는 벤더 모델 카드 또는 공식 블로그, 집계는 systemonemodels.org 디렉터리의 서술입니다. 마지막 줄의 Jev는 베이스, 크기, 헤드 구조 모두 비공개입니다. 공식 서술은 새 아키텍처, 병렬 샘플러, RLCD 학습이라는 세 가지뿐입니다.")
meta("모델별 구조 — 확인된 만큼만", "4열 표(모델·구조·판독·근거). Jev 행은 붉은색으로 '비공개' 강조.", ["Clef: joint schema head", "Kev·Strands: LoRA + 포인터 헤드", "Jev: 비공개"],
     "근거 열의 '집계'는 제3자 디렉터리의 서술로 벤더 1차 확인이 필요하다.")

# ---- 15. Jev 내부
s = base("Jev 내부 — 공개된 것과 아닌 것")
box(s, 0.7, 1.5, 5.8, 5.2, "EEF7EE", GREEN)
txt(s, 0.9, 1.6, 5.4, 0.5, "공식 진술 (3문장)", 20, True, "22502A")
txt(s, 0.9, 2.3, 5.4, 3.5, ["‘새 모델 아키텍처 + 병렬 샘플러’", "‘모든 출력이 하나의 쿼리에서 생성’", "RLCD: 교정된 결정을 위한 강화학습", "", "컨텍스트 64k(질문당 32k) · $0.042/MTok", "1.14: box 타입 추가"], 18, False, DARK)
box(s, 6.8, 1.5, 5.8, 5.2, "FDF1EE", RED)
txt(s, 7.0, 1.6, 5.4, 0.5, "비공개 · 미확인", 20, True, "6B2A1A")
txt(s, 7.0, 2.3, 5.4, 3.5, ["베이스 모델 · 크기 · 헤드 구조", "학습 데이터", "output_tokens ≈ 질문당 20의 의미", "어휘 판독인지 헤드형인지", "", "정황: 첫 선택지 편향 인정,", "score 레벨 독립 판단 서술"], 18, False, DARK)
notes(s, "Jev의 공식 진술은 세 문장입니다. 새 아키텍처와 병렬 샘플러, 모든 출력이 하나의 쿼리에서 생성된다는 것, 그리고 RLCD 학습입니다. 베이스 모델, 크기, 헤드 구조, 학습 데이터는 비공개입니다. 첫 선택지 편향을 인정한 문서와 score 레벨을 독립적으로 판단한다는 서술은 정황일 뿐이며 어휘 판독인지 헤드형인지를 단정할 수 없습니다. Cloudflare Clef의 학습 설명에도 RLCD라는 이름이 나오므로, 특정 회사의 기법이라기보다 교정을 보상으로 삼는 강화학습의 통칭일 가능성이 있습니다.")
meta("Jev 내부 — 공개된 것과 아닌 것", "좌우 2분할: 녹색(공식 진술), 붉은색(비공개·미확인).", ["공식 서술은 3문장", "베이스·헤드·데이터 비공개", "정황만 존재"],
     "Clef 학습 설명에도 RLCD가 등장하므로 특정 회사의 기법이라기보다 통칭일 가능성이 있다.")

# ---- 16. 직접 돌리기
s = base("직접 돌리기", "호스팅은 JSON 그대로, 자체 호스팅은 요건이 제각각")
rows = [["모델", "요건 · 실행", "근거"],
        ["d1-3B", "transformers ≥ 5.14, bf16 · RTX 4090 8 ms · Orin Nano 50 ms", "카드·블로그"],
        ["d1-omni-600M", "transformers ≥ 5.15, fp16 · ctx 16,384", "카드"],
        ["GLiNER2.5-Decide", "pip install gliner2 · CPU/GPU · 340M", "카드"],
        ["Clef 27B", "단일 H200에서 테스트 · SGLang /v1/systemone · Docker", "카드"],
        ["pplx-decider v1.2 / flash", "약 52 GiB / 약 9 GiB (맥 앱 16 GB)", "집계"],
        ["GLiDE", "bf16 52.2 GB · 80 GB급 GPU 1장", "집계"],
        ["Kev", "0.8–9B 맥(MLX) 가능 · 27B는 80 GB · 요청 직렬 처리", "집계"],
        ["Von", "CPU(OpenVINO) 가능 · ctx 8,192 · 유지보수 중단", "집계"]]
table(s, 0.6, 1.6, 7.8, rows, [1.8, 4.6, 1.1], 12, 0.58)
code(s, 8.65, 1.6, 4.1, 3.2, ["model = AutoModel.from_pretrained(", '  "LiquidAI/d1-3B",', "  trust_remote_code=True,", "  dtype=torch.bfloat16).to('cuda')", "", "model.system_one(text, questions)", "model.system_one_batch(", "  [(t, qs) for t in tickets])"], 11)
card(s, 8.65, 5.0, 4.1, 1.75, "교체는 설정 한 줄", "다수가 /v1/systemone 호환\nbase_url만 바꾸면 전환", PALE, MID, BLUE, 15, 13)
notes(s, "실행 방법입니다. 호스팅 API는 앞서 본 JSON을 그대로 보내면 됩니다. 자체 호스팅은 요건이 제각각입니다. d1-3B는 transformers 5.14 이상에서 몇 줄로 로드하고, system_one_batch가 여러 요청을 패딩 없이 패킹합니다. GPU에서는 새 입력 shape의 첫 호출에 오버헤드가 있으므로 워밍업을 권합니다. 다수 모델이 Jev 호환 서버를 제공하므로 로컬과 호스팅의 교체는 base_url 변경으로 가능합니다.")
meta("직접 돌리기", "좌: 요건 표(8모델). 우상: d1-3B 로컬 코드. 우하: '교체는 설정 한 줄' 카드.", ["d1-3B: 몇 줄로 로컬 실행", "27B급은 52~80 GB GPU", "호환 서버로 교체 용이"],
     "GPU에서 새 입력 shape의 첫 호출에는 오버헤드가 있으므로 서빙할 shape을 워밍업한다.")

# ---- 17. DIV
s = divider("03", "선택지 순서 의존성", "Consider, Then Choose — Runoff")
notes(s, "세 번째 묶음은 선택지 순서 의존성입니다.")
meta("03 선택지 순서 의존성", "구간 표지.", ["Runoff"], "순서 의존성 구간의 시작.")

# ---- 18. 순서만 바꿔도
s = base("순서만 바꿔도 답이 바뀐다", "flip = 선택지를 한 번 섞었을 때 답이 바뀌는 비율 (Runoff 논문)")
bars = [("Qwen3-8B listwise K=2", 4.7, MID), ("Qwen3-8B listwise K=20", 17.0, RED), ("Qwen3-8B AnyJev K=20", 10.3, AMBER),
        ("Jev 1.13 native (5셋 평균)", 5.3, MID), ("GPT-6 Luna native", 13.2, RED), ("Mercury Decide native", 9.8, AMBER)]
for i, (l, v, c) in enumerate(bars):
    y = 1.6 + i * 0.74
    txt(s, 0.7, y, 4.6, 0.55, l, 14, False, DARK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 5.4, y + 0.08, 6.0 * v / 20, 0.42, c, None, MSO_SHAPE.RECTANGLE)
    txt(s, 5.5 + 6.0 * v / 20, y, 1.2, 0.55, f"{v}%", 14, True, c, anchor=MSO_ANCHOR.MIDDLE)
card(s, 0.7, 6.3, 12.0, 0.65, "원인: 모델이 상위 2~3개 사이에서 망설일 때 flip이 집중 (pointwise 1위 확률 < 0.8이면 flip > 1/3)", None, "FDF6E8", AMBER, "7A5710", 14)
notes(s, "선택지 개수 K가 늘수록 listwise 판독의 flip이 커집니다. Qwen3-8B는 K가 2에서 20으로 갈 때 4.7퍼센트에서 17퍼센트가 됩니다. Jev 1.13은 다섯 데이터셋 평균 5.3퍼센트, GPT-6 Luna는 13.2퍼센트입니다. 원인은 모델이 상위 몇 개 후보 사이에서 망설일 때 뒤집힘이 집중된다는 점입니다. Jev 문서도 1.13이 첫 번째 선택지 쪽으로 기운다고 인정합니다.")
meta("순서만 바꿔도 답이 바뀐다", "가로 막대 6개(flip 비율). 하단에 원인 요약 카드.", ["K=20에서 17%", "Jev 1.13 평균 5.3%", "상위 후보 사이 망설임이 원인"],
     "Jev 문서 약점 8번이 같은 증상을 인정한다. 순서 편향은 선택지 이름 편향과는 별개다.")

# ---- 19. Runoff 흐름
s = base("Runoff — consider → choose", "훈련 없는 두 단계 · 순서 불변이 구조적으로 보장")
box(s, 0.6, 1.6, 5.9, 4.6, "EEF4FB", MID)
txt(s, 0.8, 1.65, 5.5, 0.4, "① Consider (pointwise · 순서 무관)", 16, True, BLUE)
chip(s, 0.9, 2.3, 5.3, 0.7, "선택지마다 따로 yes/no  f = log P(yes) − log P(no)", WHITE, BLUE, 13)
arrow(s, 3.55, 3.02, 3.55, 3.35, GREY)
chip(s, 0.9, 3.4, 5.3, 0.7, "softmax → 분포 π", WHITE, BLUE, 13)
arrow(s, 3.55, 4.12, 3.55, 4.45, GREY)
chip(s, 0.9, 4.5, 5.3, 1.0, ["질량 ≥ 0.99, 크기 ≥ 3인", "최소 상위 집합 C"], WHITE, BLUE, 13)
arrow(s, 6.55, 4.0, 7.0, 4.0, AMBER, 3)
box(s, 7.05, 1.6, 5.7, 4.6, "F3EEF9", PURPLE)
txt(s, 7.25, 1.65, 5.3, 0.4, "② Choose (listwise · C 안에서만)", 16, True, "3D2A55")
chip(s, 7.35, 2.3, 5.1, 0.7, "π 순 정렬 = canonical ranking", WHITE, "3D2A55", 13)
arrow(s, 9.9, 3.02, 9.9, 3.35, GREY)
chip(s, 7.35, 3.4, 5.1, 0.7, "m개 순환 이동마다 listwise 1회", WHITE, "3D2A55", 13)
arrow(s, 9.9, 4.12, 9.9, 4.45, GREY)
chip(s, 7.35, 4.5, 5.1, 1.0, ["로그 확률 평균 → r", "최종: C 밖 = π, C 안 = π(C)·r"], WHITE, "3D2A55", 13)
txt(s, 0.6, 6.35, 12.1, 0.5, "비용 K + m² 슬롯 (AnyJev는 K²) · 두 번의 배치 forward pass · 결정성 가정 하에서 순서 불변", 15, True, NAVY)
notes(s, "Runoff는 두 단계입니다. 고려 단계에서는 선택지를 하나씩 따로 yes 또는 no로 평가해 순서와 무관한 분포를 얻고, 질량이 0.99 이상이 되는 최소 상위 집합을 고릅니다. 선택 단계에서는 그 집합을 분포 순으로 정렬한 뒤 모든 순환 이동에 대해 listwise 판독을 하고 로그 확률을 평균합니다. 요청이 후보 집합과 분포에만 의존하므로 입력 순서가 바뀌어도 같은 결과가 나옵니다. 비용은 K 더하기 m 제곱입니다.")
meta("Runoff — consider → choose", "좌: Consider 박스(pointwise yes/no → softmax π → 상위 집합 C). 우: Choose 박스(정렬 → 순환 이동 listwise → 평균). 하단 비용 문구.",
     ["① 순서 무관하게 후보 추리기", "② 후보만 순환시켜 선택", "비용 K+m²"],
     "결정성 가정이 깨지는 닫힌 API에서는 소수의 flip이 남는다.")

# ---- 20. Runoff 결과
s = base("Runoff 결과", "정확도 % / flip % · 5개 데이터셋 평균")
rows = [["모델", "방식", "정확도", "flip"],
        ["Qwen3-8B", "Listwise → Runoff", "63.5 → 65.8", "22.8 → 0.0"],
        ["Jev 1.13", "Native → Runoff", "80.6 → 81.5", "5.3 → 0.9"],
        ["GPT-6 Luna", "Native → Runoff", "75.4 → 77.0", "13.2 → 0.0"],
        ["Mercury Decide", "Native → Runoff", "76.7 → 78.5", "9.8 → 1.4"]]
table(s, 0.6, 1.6, 7.3, rows, [1.8, 2.0, 1.8, 1.7], 14, 0.62, hl={(1, 3): GREEN, (2, 3): GREEN, (3, 3): GREEN, (4, 3): GREEN})
card(s, 8.2, 1.6, 4.5, 1.7, "슬롯: K=20, m=3", "AnyJev 400  vs  Runoff 29", PALE, MID, BLUE, 16, 18)
card(s, 8.2, 3.5, 4.5, 1.6, "지연 (Qwen3-8B, K=20)", "308 ms → 77 ms (4.0배)", PALE, MID, BLUE, 16, 18)
card(s, 0.6, 5.0, 7.3, 1.75, "한계", "선택지 4~6개면 이점 작음(0.9배) · 닫힌 API 잔존 flip\n순서만 해결, 이름 선호 편향은 별개 · MMLU-Pro에서 AnyJev에 몇 점 뒤짐", "FDF1EE", RED, "6B2A1A", 16, 13)
card(s, 8.2, 5.3, 4.5, 1.45, "Jev에서 구현", "K개 noul 1요청 + choice m개 요청", "EEF7EE", GREEN, "22502A", 15, 14)
notes(s, "결과입니다. 여덟 개 시스템 원 모델에 씌웠을 때 순서 의존성이 사실상 사라졌고 정확도는 같거나 올랐습니다. 닫힌 API의 잔존 flip은 같은 요청에도 확률이 달라지는 비결정성 때문이라고 논문이 설명합니다. K가 20, 후보가 3이면 슬롯이 400에서 29로 줄어듭니다. 선택지가 4에서 6개로 적으면 이점이 작습니다. Jev API에서도 noul을 선택지마다 한 요청에 묶어 consider 단계를, choice 요청 m개로 choose 단계를 구현할 수 있습니다.")
meta("Runoff 결과", "좌: 결과 표(4모델, 정확도·flip). 우: 슬롯 비교·지연 카드. 하단: 한계와 Jev 구현 카드.", ["flip 0~1.4%", "슬롯 400 → 29", "4배 빠름 (K=20)"],
     "선택지가 적으면 이점이 작다. 이 글의 Jev 어댑터 코드는 문서 사양 기준 스케치이며 실제 호출로 검증하지 않았다.")

# ---- 21. DIV
s = divider("04", "지형과 성능", "누가 무엇을 냈고, 숫자는 얼마나 믿을 수 있나")
notes(s, "네 번째 묶음은 모델 지형과 성능입니다.")
meta("04 지형과 성능", "구간 표지.", ["모델 지형", "벤치마크 · 비용 · 지연"], "지형과 성능 구간의 시작.")

# ---- 22. 모델 지형
s = base("모델 지형", "가격은 입력 $/MTok · 출력은 모두 무료 · 대부분 집계 사이트 표기")
rows = [["모델", "공개", "백본 · 규모", "가격", "가중치"],
        ["Jev 1.13/1.14", "09-15 / 10-14", "비공개", "0.042", "닫힘"],
        ["OpenAI Decisions", "10-06 베타", "gpt-6-luna", "0.10", "닫힘"],
        ["Liquid d1 / d1-ultra", "09-29 / 10-15", "비공개 / 더 큰 오픈 모델", "0.02 / 0.042", "닫힘"],
        ["Liquid d1-3B · omni-600M", "10-07", "VLM 3.12B · 인코더 587M", "로컬", "LFM Open"],
        ["Cloudflare Clef / flash", "10-01", "Qwen3.8-27B / 3.5-9B", "집계 0.24", "Apache-2.0"],
        ["pplx-decider v1.2", "10-01 (v1.2 10-16)", "Qwen3.8-27B / 4B 증류", "0.02 / 0.01", "Apache-2.0"],
        ["Fastino GLiDE 1.2", "09-30 / 10-22", "Qwen3.8-27B", "0.04", "Community"],
        ["GLiNER2.5-Decide", "09-24", "DeBERTa-v3-large 340M", "로컬", "Apache-2.0"],
        ["Span-02 (Respan)", "10-15", "Gemma 4 26B-A4B", "0.01", "Apache-2.0"],
        ["Kev · Strands · Von · CLM", "9~10월", "LoRA+헤드 · 인코더 · 대조", "로컬", "Apache-2.0"]]
table(s, 0.5, 1.55, 12.3, rows, [2.7, 1.8, 3.2, 1.4, 1.6], 12, 0.47)
txt(s, 0.5, 6.7, 12.3, 0.4, "Microsoft-Decision-1 · Solar Decide · Mercury Decide · Tev1 · Clairic 등은 본문 표 참조", 12, False, GREY)
notes(s, "대표 모델만 추렸습니다. 가격과 공개일은 대부분 집계 사이트 표기이므로 도입 전에 벤더 가격표로 재확인해야 합니다. 세 가지가 눈에 띕니다. 공개 구조의 대부분이 Qwen이나 Gemma 위의 후처리라는 백본 수렴, 340M에서 27B까지의 크기 스펙트럼, 그리고 한 달에 버전이 세 번 바뀌는 속도입니다. 별칭을 쓰면 튜닝한 임계값이 어긋나므로 버전 ID를 고정해야 합니다.")
meta("모델 지형", "10행 표(모델·공개·백본·가격·가중치).", ["백본 수렴 (Qwen · Gemma)", "340M ~ 27B 스펙트럼", "버전이 한 달에 3번"],
     "pplx-decider는 10월 한 달에 v1, v1.1, v1.2로 바뀌었다. 별칭이 움직이면 임계값이 어긋나므로 버전 ID를 고정한다.")

# ---- 23. AA 지수
s = base("Artificial Analysis Decision Index", "2026-10-13 · 지식 · 추론 · 도구 호출 · 안전 분류")
data = [("Claude Haiku 5.5 (LM)", 61, "$184", "1.4 s", GREY), ("Jev 1.13", 58, "$42", "0.13 s", AMBER), ("GPT-6 Luna (LM)", 55, "$131", "1.5 s", GREY),
        ("GLiDE (런치판)", 54, "$119", "—", MID), ("Microsoft-Decision-1", 54, "$31", "—", MID)]
txt(s, 5.6, 1.5, 1.2, 0.3, "지수", 12, True, GREY)
txt(s, 9.2, 1.5, 1.5, 0.3, "$ / 100만 과제", 12, True, GREY)
txt(s, 11.2, 1.5, 1.5, 0.3, "중앙 지연", 12, True, GREY)
for i, (l, v, c, t, col) in enumerate(data):
    y = 1.9 + i * 0.72
    txt(s, 0.6, y, 3.6, 0.55, l, 14, False, DARK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 4.0, y + 0.08, 4.8 * v / 65, 0.4, col, None, MSO_SHAPE.RECTANGLE)
    txt(s, 4.05 + 4.8 * v / 65, y, 0.8, 0.55, str(v), 14, True, col, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 9.2, y, 1.5, 0.55, c, 14, False, DARK, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 11.2, y, 1.5, 0.55, t, 14, False, DARK, anchor=MSO_ANCHOR.MIDDLE)
card(s, 0.6, 5.6, 4.0, 1.25, "강한 곳", "분류 · 도구 호출 판단", "EEF7EE", GREEN, "22502A", 15, 14)
card(s, 4.75, 5.6, 4.0, 1.25, "약한 곳", "지식 · 추론 (AA-Omni 35 vs 52~66)", "FDF1EE", RED, "6B2A1A", 15, 14)
card(s, 8.9, 5.6, 3.8, 1.25, "과신", "거의 모든 구간에서 정확도 < 제시 확률", "FDF6E8", AMBER, "7A5710", 15, 14)
notes(s, "AA 지수입니다. Jev 1.13은 결정 모델 중 선두권인 58점이고, 가장 높은 언어 모델 Haiku 5.5의 61점에 3점 뒤지지만 비용은 약 4.4분의 1, 지연은 약 11분의 1입니다. 세부로 보면 분류와 도구 호출 판단에서는 결정 모델이 앞서거나 비슷하고 지식과 추론에서는 뒤집니다. 그리고 AA는 거의 모든 확률 구간에서 모든 결정 모델의 정확도가 제시 확률보다 낮다고 보고합니다. 확률이 붙는다는 것과 그 확률을 믿어도 된다는 것은 다른 문제입니다.")
meta("Artificial Analysis Decision Index", "5행 가로 막대(지수) + 비용·지연 열. 하단 강한 곳·약한 곳·과신 카드.", ["Jev 58 vs Haiku 61", "비용 1/4.4 · 지연 1/11", "과신"],
     "AA는 결정 모델 전부가 과신한다고 보고한다. LLM 쪽 교정은 발표 시점에 coming soon이었다.")

# ---- 24. 지표 불일치
s = base("숫자가 서로 어긋나는 이유", "‘1위’는 지표 이름과 함께가 아니면 의미가 없다")
cs = [("지표마다 순위가 다르다", "AA: Jev 1.13 = 결정 모델 선두권\nJev Decision Index 0.4.1: GLiDE 71.5, pplx 68.0, Jev 1.14 60.2"),
      ("프로토콜이 다르다", "BenchLM의 d1-ultra 65.2는\nAA의 Jev 58과 같은 지수가 아님"),
      ("지연도 측정마다 다르다", "Jev 중앙값: AA 130 ms\nCloudflare 측정 524 ms\n(회사 주장 70–500 ms)")]
for i, (h, d) in enumerate(cs):
    card(s, 0.6 + i * 4.1, 1.7, 3.9, 3.2, h, d.split("\n"), PALE, MID, BLUE, 18, 16)
rows = [["ECE (낮을수록 좋음)", "GLiDE", "Kev-27B", "pplx v1.2", "Jev 1.14", "Von", "CLM-8B"],
        ["Jev Decision Index 0.4.1", "0.020", "0.022", "0.055", "0.080", "0.216", "0.323"]]
table(s, 0.6, 5.2, 12.1, rows, [2.6, 1.3, 1.3, 1.4, 1.3, 1.3, 1.3], 13, 0.55)
txt(s, 0.6, 6.45, 12.1, 0.5, "서드파티 리더보드의 비공개 테스트 — 자기 도메인에서 신뢰도 곡선을 직접 그려야 한다", 14, True, RED)
notes(s, "숫자가 어긋나는 이유를 정리합니다. 첫째, 같은 모델이 지표에 따라 순위가 크게 달라집니다. 둘째, BenchLM이 표기한 d1-ultra의 65.2는 다른 프로토콜의 지수라 AA의 Jev 58과 비교할 수 없습니다. 셋째, 같은 Jev의 중앙 지연이 AA에서는 130밀리초, Cloudflare 측정에서는 524밀리초입니다. 측정 지점이 다르기 때문입니다. 교정 오차는 오픈 가중치 소형 모델이 낮은 경우도 있지만 서드파티 비공개 테스트 결과이므로 자기 도메인에서 확인해야 합니다.")
meta("숫자가 서로 어긋나는 이유", "3 카드(순위·프로토콜·지연 불일치) + ECE 비교 표.", ["지표마다 순위가 다르다", "프로토콜이 다르다", "지연은 측정 지점마다 다르다"],
     "순위 비교에는 단일 출처 안의 값만 쓴다.")

# ---- 25. DIV
s = divider("05", "실전 패턴", "게이팅 · 캐스케이드 · 팬아웃 · UI 라우터")
notes(s, "다섯 번째 묶음은 실전 패턴입니다.")
meta("05 실전 패턴", "구간 표지.", ["게이팅 · 캐스케이드 · 팬아웃 · UI 라우터"], "실전 패턴 구간의 시작.")

# ---- 26. 게이팅·캐스케이드
s = base("신뢰도 게이팅 + 캐스케이드", "싼 모델로 만들고 · 결정 모델로 검증하고 · 의심스러울 때만 비싼 모델로")
chip(s, 0.6, 2.0, 1.9, 0.9, ["문서·요청"], PALE, BLUE, 14)
arrow(s, 2.52, 2.45, 2.9, 2.45)
chip(s, 2.95, 2.0, 2.3, 0.9, ["① 저가 LLM 추출", "gpt-5.4-mini"], PALE, BLUE, 13)
arrow(s, 5.27, 2.45, 5.65, 2.45)
chip(s, 5.7, 2.0, 2.6, 0.9, ["② System One 검증", "필드별 noul 배터리"], AMBER, NAVY, 13)
arrow(s, 8.32, 2.45, 8.7, 2.45)
chip(s, 8.75, 2.0, 1.9, 0.9, ["③ any_flag", "P > 0.7 ?"], WHITE, BLUE, 13)
arrow(s, 10.65, 2.2, 11.05, 1.75, GREEN)
chip(s, 11.1, 1.45, 1.8, 0.65, "저가 결과 채택", "EEF7EE", "22502A", 12)
arrow(s, 10.65, 2.7, 11.05, 3.15, RED)
chip(s, 11.1, 2.9, 1.8, 0.8, ["④ 추론 모델", "gpt-5.5 high"], "FDF1EE", "6B2A1A", 12)
card(s, 0.6, 3.9, 5.9, 2.8, "confidence 3구간", ["높음 (>0.9)  자동 실행", "중간  사용자 확인", "낮음 (<0.5)  사람·LLM 폴백", "noul 시작점 0.7 / 0.3 — 위험에 비례"], PALE, MID, BLUE, 17, 16)
card(s, 6.8, 3.9, 5.9, 2.8, "SDE 쿡북 (TypeSafe 내부 100건)", ["추론 모델 단독: 품질 ~0.81, 건당 ~$0.10", "캐스케이드 프런티어가 모든 단일 모델의 위·왼쪽", "검증 질문: 환각 · 형식 위반 · 불완전 …", "정확한 좌표는 문서에 없음"], PALE, MID, BLUE, 17, 16)
notes(s, "가장 정석적인 패턴입니다. 저가 LLM이 추출하고, 결정 모델이 필드별로 문제 있음의 확률을 평가하고, 임계값 0.7을 하나라도 넘으면 추론 모델로 올립니다. TypeSafe 쿡북의 내부 100건 평가에서 이 캐스케이드의 비용 대비 품질 프런티어는 모든 단일 모델의 위쪽에 놓였습니다. 다만 정확한 달러와 품질 좌표는 문서에 없습니다. 임계값은 위험에 비례하게 정하고, 자기 데이터로 신뢰도 곡선을 그려 확정해야 합니다.")
meta("신뢰도 게이팅 + 캐스케이드", "상단: 요청→저가 LLM→System One 검증→any_flag 분기(채택/추론 모델 에스컬레이션). 하단 2카드.", ["confidence 3구간", "필드별 noul 배터리", "FIRE_T = 0.7"],
     "SDE 쿡북은 내부 평가이며 비용은 현재 Jev 단가로 재계산되지 않았다.")

# ---- 27. 패턴 4종
s = base("그 밖의 패턴", "결정 모델은 ‘고르고 점수 매기는’ 역할만")
pc = [("팬아웃", "13개 질문을 한 번에", "12.2배 저렴 · 10.0배 빠름\n(질문은 독립 평가)"),
      ("재정렬", "후보마다 noul 1개", "top-1 5% → 18%\n1,200호출 $0.0645"),
      ("UI 라우터", "CopilotKit · json-render", "choice(control) + score(fit_)\n앱이 검증 · 실행은 코드"),
      ("도메인 특화", "CS1-4B 임상", "14과제 중 13 승\nLM 토큰 690배 절감")]
for i, (h, s1, d) in enumerate(pc):
    x = 0.6 + i * 3.1
    box(s, x, 1.7, 2.95, 4.2, PALE, MID)
    txt(s, x + 0.15, 1.8, 2.7, 0.5, h, 21, True, BLUE)
    txt(s, x + 0.15, 2.45, 2.7, 0.8, s1, 14, True, AMBER)
    txt(s, x + 0.15, 3.4, 2.7, 2.3, d.split("\n"), 18, False, DARK)
txt(s, 0.6, 6.1, 12.1, 0.8, "원칙: 후보 정의 · 검증 · 실행 권한은 코드가 쥔다 — 결정 모델은 앱이 정의한 후보 중에서 고르기만 한다", 16, True, NAVY)
notes(s, "나머지 패턴입니다. 팬아웃은 질문이 독립 평가이므로 한 요청에 몰아넣으면 쿡북 기준 12.2배 저렴하고 10배 빠릅니다. 재정렬은 후보마다 noul 하나를 물어 상위 30건을 재정렬해 top-1 적중률을 5퍼센트에서 18퍼센트로 올렸습니다. UI 라우터 사례에서는 Jev가 후보 중 고르고 점수만 매기며 검증과 실행은 앱이 합니다. 임상 기록용 CS1-4B는 작은 결정 모델이 대량의 기록을 읽고 범용 LM은 질문과 종합을 맡는 분업을 보여 줍니다. OpenIntelligentUI가 이 구조를 매 턴 렌더러 선택에 쓰는 것은 자매 글에서 다룹니다.")
meta("그 밖의 패턴", "4 카드(팬아웃·재정렬·UI 라우터·도메인 특화), 하단 원칙 문장.", ["팬아웃 12.2배", "재정렬 5%→18%", "UI 라우터", "CS1 LM 토큰 690배 절감"],
     "자매 글에서 OpenIntelligentUI가 매 턴 렌더러 선택에 Jev를 쓰는 사례를 다룬다.")

# ---- 28. 맡기지 말 것
s = base("맡기지 말아야 할 일", "Jev 1.13의 알려진 약점 — 코드로 풀어라")
rows = [["약점", "회피법"],
        ["개수 세기 · 수치 비교 · 16진수", "후보마다 질문 1개, 합산은 코드"],
        ["날짜 · 기간 계산", "월·일·연을 choice로 추출 후 코드 산술"],
        ["긴 state에 섞인 무관 정보", "코드로 거르거나 noul로 필터"],
        ["적대적 텍스트 주입", "기준 명시 + 에지 케이스 사전 테스트"],
        ["첫 선택지로 기움", "순서 교란 검사 · Runoff"],
        ["score 수치 보간", "임계값 통과 여부에만 사용"],
        ["텍스트 생성", "후보는 정규식·생성 모델이, Jev는 고르기만"]]
table(s, 0.6, 1.55, 12.1, rows, [4.5, 7.0], 15, 0.58)
txt(s, 0.6, 6.3, 12.1, 0.4, "결정 모델은 의미 판단의 전문가이지 계산기가 아니다", 16, True, NAVY)
notes(s, "TypeSafe가 스스로 공개한 약점 목록이 곧 실무 체크리스트입니다. 개수 세기, 수치 비교, 날짜 계산처럼 코드가 정확히 계산할 수 있는 일은 코드로 처리해야 합니다. 관련 없는 긴 state는 분산 요인이 되고, 적대적 텍스트에 취약하며, 첫 선택지로 기우는 경향도 있습니다. 결정 모델은 의미 판단의 전문가이지 계산기가 아닙니다.")
meta("맡기지 말아야 할 일", "2열 표(약점·회피법) 7행.", ["계산·날짜는 코드로", "무관한 state는 걸러서", "점수는 임계값 통과만"],
     "출처는 TypeSafe의 Jev 1.13 jaggedness 문서(2026-10-02 검토).")

# ---- 29. 모르는 것
s = base("아직 모르는 것")
unk = [("Jev 내부", "베이스·헤드·데이터 · output_tokens의 의미"), ("OpenAI Decisions", "내부 방식 · ‘10배 빠름’의 벤치마크 · GA 시점"),
       ("Microsoft-Decision-1", "블로그 본문 미확인 · 프리뷰/GA 상충"), ("Tev1 · Mercury", "1차 자료 미확인 · 아키텍처 비공개"),
       ("LLM 교정 비교", "AA 발표 시점 coming soon"), ("가격의 지속가능성", "보조금 가능성 인정 · 한 달 새 단가 반토막")]
for i, (h, d) in enumerate(unk):
    cx = 0.6 + (i % 2) * 6.15
    cy = 1.45 + (i // 2) * 1.65
    card(s, cx, cy, 5.95, 1.45, h, d, "FDF1EE" if i < 2 else PALE, RED if i < 2 else MID, "6B2A1A" if i < 2 else BLUE, 17, 14)
txt(s, 0.6, 6.45, 12, 0.4, "가장 큰 구조적 미지수: ‘출력 토큰 0’ 서술과 usage.output_tokens의 불일치를 벤더가 설명한 적이 없다", 13, True, RED)
notes(s, "아직 모르는 것들입니다. Jev의 내부는 공개되지 않았습니다. OpenAI Decisions의 내부 방식과 속도 주장의 근거도 서술되어 있지 않습니다. Microsoft-Decision-1은 블로그 제목만 확인했습니다. 가장 큰 구조적 미지수는 출력 토큰 0이라는 서술과 usage의 output_tokens 필드가 서로 맞지 않는 점을 벤더가 설명한 적이 없다는 것입니다.")
meta("아직 모르는 것", "2×3 카드 그리드(붉은색 2장은 핵심 미지수).", ["Jev 내부", "OpenAI 내부 방식", "독립 벤치마크 부재"],
     "벤더 주장 중 독립 검증이 없는 수치는 모두 '주장'으로 표기했다.")

# ---- 30. 결론
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
box(s, 0, 0, 0.35, 7.5, AMBER, None, MSO_SHAPE.RECTANGLE)
txt(s, 0.9, 0.8, 11.5, 0.9, "정리 · Q&A", 40, True, WHITE)
pts = ["계약: state + questions → probabilities · confidence. 호환 구현이 표준", "실행: prefill 1회 + 판독. 단위와 방식은 벤더마다 다르다", "약점: 선택지 순서(Runoff)와 과신(임계값은 자기 데이터로)", "패턴: 싼 생성 → 결정 모델 검증 → 의심될 때만 비싼 모델", "확인된 것과 집계된 것을 구분해서 읽자"]
for i, p in enumerate(pts):
    txt(s, 0.9, 2.0 + i * 0.85, 11.6, 0.7, "•  " + p, 20, False, "E6EDF5")
txt(s, 0.9, 6.5, 11, 0.5, "자매 글: OpenIntelligentUI — Jev를 매 턴 렌더러 라우터로 쓰는 사례", 14, False, "9FB6D0")
notes(s, "정리합니다. 계약은 수렴했고, 실행 방식은 벤더마다 다르며, 약점은 순서 의존성과 과신입니다. 실전에서는 싼 모델이 생성하고 결정 모델이 검증하며 의심스러울 때만 비싼 모델을 부르는 구조가 정석입니다. 마지막으로, 확인된 사실과 집계된 전언을 구분해서 읽으시기를 권합니다. 질문을 받겠습니다.")
meta("정리 · Q&A", "남색 배경, 5개 요약 문장.", ["계약 수렴", "실행 방식은 제각각", "순서 · 과신이 약점", "싼 생성 → 검증 → 에스컬레이션"],
     "자매 글 OpenIntelligentUI에서 Jev를 렌더러 라우터로 쓰는 사례를 다룬다.")

# ================================================================ 저장
OUT = "jev-decision-models-presentation.pptx"
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))

# ---- 아웃라인 MD
md = ["# Jev와 System One 모델 — 발표 아웃라인", "",
      "> 📊 **발표자료**: [jev-decision-models-presentation.pptx](./jev-decision-models-presentation.pptx)", "", "---", ""]
for i, m in enumerate(SLIDES, 1):
    md += [f"## Slide {i}: {m['title']}", "", f"**Visual**: {m['visual']}", "", "**Key Points**:"]
    md += [f"- {p}" for p in m["points"]]
    md += ["", f"**Speaker Notes**: {m['notes']}", "", "---", ""]
open("jev-decision-models-presentation.md", "w").write("\n".join(md))
print("outline md written:", len(SLIDES))
