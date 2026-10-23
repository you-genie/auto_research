"""
OpenIntelligentUI 서베이 — 참고문헌 XLSX 생성 스크립트
"""
import os

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


references = [
    # 원본 제품
    {
        "번호": 1,
        "제목": "GPT-6 and Intelligent UI for everyone",
        "저자/출처": "OpenAI",
        "URL": "https://openai.com/index/gpt-6-for-everyone/",
        "발행일": "2026-10-07",
        "요약": "ChatGPT Intelligent UI 공식 발표. 네이티브 스트리밍 컴포넌트 라이브러리와, 모델이 생성하는 동안 UI를 처리하는 컴파일러로 구성. GPT-6가 라이브러리 사용법과 상호작용/텍스트 선택 기준을 학습했다고 설명. 원문 직접 열람은 403으로 막혀 검색 결과 노출 본문 기준으로 인용.",
        "관련 섹션": "1. 문제와 이름 충돌"
    },
    {
        "번호": 2,
        "제목": "OpenAI Rolls Out GPT-6 and Intelligent UI in ChatGPT",
        "저자/출처": "Let's Data Science (2차 보도)",
        "URL": "https://letsdatascience.com/news/openai-brings-gpt-6-and-intelligent-ui-to-chatgpt-4f0d797c",
        "발행일": "2026-10",
        "요약": "롤아웃 일정(유료 플랜 우선, Free/Go 다음 날), Chat 경험 한정, GPT-6 Sol/Luna 적용 범위를 정리한 2차 보도. 1차 자료로 OpenAI 게시글을 링크.",
        "관련 섹션": "1. 문제와 이름 충돌"
    },
    # 구현 A
    {
        "번호": 3,
        "제목": "CopilotKit/OpenIntelligentUI README",
        "저자/출처": "CopilotKit (GitHub, MIT)",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI",
        "발행일": "2026-10",
        "요약": "Next.js + CopilotKit 프런트엔드, Python Deep Agent(FastAPI), 선택적 MCP 서버, 디자인 시스템 패키지로 구성된 오픈소스 구현. chat-latest 답변 모델과 Jev 라우터, generateSandboxedUi 파라미터 순서, BYOK 키 처리, 설치 절차 기술.",
        "관련 섹션": "1, 3"
    },
    {
        "번호": 4,
        "제목": "Architecture (docs/architecture.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/architecture.md",
        "발행일": "2026-10",
        "요약": "구성 요소와 데이터 흐름(LangGraphHttpAgent, create_deep_agent, BoundedMemorySaver), openGenerativeUI 미들웨어, 샌드박스 브리지(sendPrompt/openLink), 프로세스 메모리 체크포인트, 패치 API 부재, query_data 동작을 설명.",
        "관련 섹션": "3, 8"
    },
    {
        "번호": 5,
        "제목": "Visualization routing (docs/visualization-routing.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/visualization-routing.md",
        "발행일": "2026-10",
        "요약": "Jev(jev-latest)가 TypeSafe System One API로 턴마다 시각화 유형을 선택하는 방식, 입력 제한(최근 4개 메시지, 각 6,000자), 선택→렌더러 도출 표(A2UI 표 vs Open Generative UI), 실패 시 폴백 없음, A2UI 표 제약(12열/200행), 지도 경로 표기 규칙.",
        "관련 섹션": "3, 4, 8"
    },
    {
        "번호": 6,
        "제목": "Generative UI (docs/generative-ui.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/generative-ui.md",
        "발행일": "2026-10",
        "요약": "useComponent 등 훅 기반 생성형 UI와 Open Generative UI 레일. generateSandboxedUi 인자 순서와 그 이유, OpenGenerativeUIMiddleware, Idiomorph 프리뷰, websandbox 부팅, 사전 주입 라이브러리(three/gsap/d3/chart.js), CSS 변수 테마 계약.",
        "관련 섹션": "3"
    },
    {
        "번호": 7,
        "제목": "Product and agent behavior (docs/interactive-answers.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/interactive-answers.md",
        "발행일": "2026-10",
        "요약": "응답 유형 선택 기준, 출처(provenance) 라벨링 규칙, 스냅샷 개념, sendPrompt 자동 전송 금지 지침, 저장소 개명 안내. 환각 완화를 프롬프트 규칙으로 다루는 방식의 근거.",
        "관련 섹션": "1, 3, 8"
    },
    {
        "번호": 8,
        "제목": "Deployment (docs/deployment.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/deployment.md",
        "발행일": "2026-10",
        "요약": "Render 블루프린트, 환경 변수(LLM_MODEL, TYPESAFE_API_KEY, JEV_MODEL), 방문자 BYOK 헤더 처리(x-openai-api-key, x-jev-api-key), 체크포인트의 프로세스 메모리 한계, 출시 검증 절차.",
        "관련 섹션": "3, 8"
    },
    {
        "번호": 9,
        "제목": "Telemetry (docs/telemetry.md)",
        "저자/출처": "CopilotKit/OpenIntelligentUI",
        "URL": "https://github.com/CopilotKit/OpenIntelligentUI/blob/main/docs/telemetry.md",
        "발행일": "2026-10",
        "요약": "설치당 최대 2개의 익명 이벤트 전송 내용과 DO_NOT_TRACK 등 옵트아웃 방법. 프롬프트·키·모델명 미전송 명시.",
        "관련 섹션": "3"
    },
    {
        "번호": 10,
        "제목": "Introducing Open Intelligent UI",
        "저자/출처": "Anmol Baranwal, Eli Berman (CopilotKit 블로그)",
        "URL": "https://www.copilotkit.ai/blog/introducing-open-intelligent-ui",
        "발행일": "2026-10-09",
        "요약": "오픈 대안 출시 글. 폐쇄성/단일 모델 한계를 동기로 제시, Jev가 포맷을 선택하고 OpenAI가 답변을 작성한다는 구성, AG-UI로 에이전트 교체 가능 주장. 홍보성 글이며 벤치마크·비용 수치 없음.",
        "관련 섹션": "1, 3, 8"
    },
    {
        "번호": 11,
        "제목": "Open Generative UI",
        "저자/출처": "CopilotKit 공식 문서",
        "URL": "https://docs.copilotkit.ai/generative-ui/open-generative-ui",
        "발행일": "2026",
        "요약": "generateSandboxedUi 단일 툴 호출, activity 이벤트 변환, 샌드박스 함수(Zod 검증), designSkill, 보안 노트(브리지 표면 최소화 권고) 설명.",
        "관련 섹션": "3"
    },
    {
        "번호": 12,
        "제목": "The Generative UI Spectrum",
        "저자/출처": "CopilotKit",
        "URL": "https://www.copilotkit.ai/generative-ui-spectrum",
        "발행일": "2026",
        "요약": "생성형 UI를 Controlled/Declarative/MCP Apps/Open으로 나눈 벤더 분류와 AG-UI·A2UI·MCP 프로토콜 스택 설명. Open 구간이 더 비싸고 덜 예측 가능하다는 평가 포함(벤더 관점).",
        "관련 섹션": "2, 7, 8, 9"
    },
    # 구현 B
    {
        "번호": 13,
        "제목": "thesysdev/open-intelligent-ui README",
        "저자/출처": "Thesys (GitHub, MIT)",
        "URL": "https://github.com/thesysdev/open-intelligent-ui",
        "발행일": "2026-10",
        "요약": "OpenUI 기반 ChatGPT Intelligent UI 재현. /api/chat, Gateway 검증·자동 수정, 직접 경로/Ollama 경로, $variables·@functions, 서버 엔드포인트, localStorage 스레드, 라이트 전용 채팅, Ollama 컨텍스트 함정, 벤치마크 스크립트 설명.",
        "관련 섹션": "1, 5, 6, 8"
    },
    {
        "번호": 14,
        "제목": "thesysdev/openui README",
        "저자/출처": "Thesys (GitHub, MIT)",
        "URL": "https://github.com/thesysdev/openui",
        "발행일": "2026",
        "요약": "OpenUI 프레임워크 개요, 패키지 구성(lang-core, 프레임워크별 런타임, a2ui), 토큰 벤치마크(7개 시나리오: 4,800 vs 10,180 vs 9,948) 및 자체 비교표(제작사 자체 측정).",
        "관련 섹션": "5, 7"
    },
    {
        "번호": 15,
        "제목": "OpenUI Lang",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs/openui-lang",
        "발행일": "2026",
        "요약": "줄 지향 문법(identifier = Expression), 위치 인자에 의한 토큰 절감 논리, 스트리밍 파싱, 잘못된 부분을 버리고 나머지를 렌더링하는 견고성.",
        "관련 섹션": "5, 7"
    },
    {
        "번호": 16,
        "제목": "OpenUI Lang Specification",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs/openui-lang/specification",
        "발행일": "2026",
        "요약": "문법 규칙(@ 접두어, 포워드 참조, 중복 이름 처리), root 진입문, 액션, 오류 코드(unknown-component 등), 스트리밍 일관성 명세.",
        "관련 섹션": "5, 7"
    },
    {
        "번호": 17,
        "제목": "OpenUI Lang Architecture",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs/openui-lang/architecture",
        "발행일": "2026",
        "요약": "생성/실행 분리, Query·Mutation·$variables·@Run·@Reset, toolProvider, 증분 편집(전체 약 400토큰 vs 패치 약 60토큰)을 코드 예제로 설명.",
        "관련 섹션": "5, 8, 9"
    },
    {
        "번호": 18,
        "제목": "Autofix",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs/autofix",
        "발행일": "2026",
        "요약": "생성 완료 후 잘못된 OpenUI Lang을 한 번의 API 호출로 수정. 오류 코드 목록, fix_summary 상태, 최대 2회 시도, 100,000자/20턴·8,000자 상한, 스트리밍 비지원.",
        "관련 섹션": "5"
    },
    {
        "번호": 19,
        "제목": "Gateway",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs/gateway",
        "발행일": "2026",
        "요약": "Chat Completions/Responses API 제공, 스트리밍 중 검증·교정, 공급자 장애 시 재시도, 공급자 가격 그대로 과금(BYOK 포함).",
        "관련 섹션": "5"
    },
    {
        "번호": 20,
        "제목": "OpenUI Docs 인덱스",
        "저자/출처": "OpenUI 공식 문서",
        "URL": "https://www.openui.com/docs",
        "발행일": "2026",
        "요약": "OpenUI를 생성형 UI 오픈 표준으로 소개하고 통합(AG-UI, LangChain, Vercel AI SDK, Google ADK, Mastra 등)과 문서 구조를 안내.",
        "관련 섹션": "2"
    },
    # Jev
    {
        "번호": 21,
        "제목": "System One API Quickstart",
        "저자/출처": "TypeSafe 공식 문서",
        "URL": "https://docs.typesafe.ai/introduction/quickstart",
        "발행일": "2026",
        "요약": "Jev(System One) API 소개. state와 questions를 받아 choice/score/noul 프리미티브별로 답과 확률·신뢰도를 반환. 엔드포인트 POST /v1/systemone, 모델명 jev-latest. 지연·가격 정보는 해당 페이지에 없음.",
        "관련 섹션": "4"
    },
    # 프로토콜
    {
        "번호": 22,
        "제목": "A2UI (Agent-to-User Interface)",
        "저자/출처": "Google / a2ui-project (GitHub, Apache-2.0)",
        "URL": "https://github.com/google/A2UI",
        "발행일": "2026",
        "요약": "에이전트가 UI 의도를 JSON으로 보내면 클라이언트가 자체 컴포넌트로 렌더링하는 선언형 포맷. 카탈로그, 평면 컴포넌트 목록, 보안 입장(코드가 아닌 데이터), 렌더러(Lit/Angular/React/Flutter), 전송(A2A/AG-UI).",
        "관련 섹션": "2, 7, 9"
    },
    {
        "번호": 23,
        "제목": "A2UI v0.9: The New Standard for Portable, Framework-Agnostic Generative UI",
        "저자/출처": "Google Developers Blog",
        "URL": "https://developers.googleblog.com/en/a2ui-v0-9-generative-ui/",
        "발행일": "2026-04-17",
        "요약": "A2UI v0.9 변경점: 공용 web core, React 렌더러, 에이전트 SDK, 클라이언트 정의 함수, 클라이언트-서버 데이터 동기화, 모듈형 스키마.",
        "관련 섹션": "2, 7"
    },
    {
        "번호": 24,
        "제목": "json-render",
        "저자/출처": "Vercel Labs (GitHub, Apache-2.0)",
        "URL": "https://github.com/vercel-labs/json-render",
        "발행일": "2026",
        "요약": "Zod 카탈로그 기반 가드레일 생성형 UI. root + elements 맵 JSON 스펙, SpecStream 스트리밍 컴파일러, 다수 렌더러, 실험적 Jev 조합 경로.",
        "관련 섹션": "2, 7"
    },
    {
        "번호": 25,
        "제목": "Vercel json-render 보도",
        "저자/출처": "InfoQ (2차 출처)",
        "URL": "https://www.infoq.com/news/2026/03/vercel-json-render",
        "발행일": "2026-03",
        "요약": "json-render 2026-01 공개와 카탈로그·Zod 구조를 다룬 보도. Vercel 측이 초기 연구 프로토타입 성격으로 설명했다는 맥락 제공.",
        "관련 섹션": "7"
    },
    {
        "번호": 26,
        "제목": "AG-UI Introduction",
        "저자/출처": "AG-UI 공식 문서",
        "URL": "https://docs.ag-ui.com/introduction",
        "발행일": "2026",
        "요약": "에이전트-사용자 앱 간 이벤트 기반 양방향 프로토콜 소개. MCP(에이전트-툴)/A2A(에이전트-에이전트)와의 층위 구분, A2UI·MCP Apps와의 관계.",
        "관련 섹션": "2"
    },
    {
        "번호": 27,
        "제목": "AG-UI Events",
        "저자/출처": "AG-UI 공식 문서",
        "URL": "https://docs.ag-ui.com/concepts/events",
        "발행일": "2026",
        "요약": "이벤트 분류: 수명주기(RunStarted 등), TextMessage*, ToolCall*, State*, Activity*, Reasoning*, Subagent*, Raw/Custom.",
        "관련 섹션": "2"
    },
    {
        "번호": 28,
        "제목": "MCP Apps Overview",
        "저자/출처": "Model Context Protocol 공식 문서",
        "URL": "https://modelcontextprotocol.io/extensions/apps/overview",
        "발행일": "2026",
        "요약": "툴 정의의 _meta.ui.resourceUri(ui://) 선언, 호스트의 샌드박스 iframe 렌더링, postMessage 기반 JSON-RPC 방언, 보안 모델과 클라이언트 지원 현황.",
        "관련 섹션": "2"
    },
]


def create_references_xlsx():
    """참고문헌 XLSX 파일을 생성합니다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "참고문헌"

    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ROW_FILL_1 = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    ROW_FILL_2 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    HEADER_FONT = Font(name="Malgun Gothic", bold=True, color="FFFFFF", size=11)
    BODY_FONT = Font(name="Malgun Gothic", size=10)
    URL_FONT = Font(name="Malgun Gothic", size=10, color="0563C1", underline="single")
    BOLD_FONT = Font(name="Malgun Gothic", bold=True, size=10)

    THIN_BORDER = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["번호", "제목", "저자/출처", "URL", "발행일", "요약 (한국어)", "관련 섹션"]
    col_widths = [6, 55, 30, 45, 12, 65, 22]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, ref in enumerate(references, 2):
        is_odd = (row_idx % 2 == 0)
        row_fill = ROW_FILL_1 if is_odd else ROW_FILL_2

        values = [
            ref["번호"],
            ref["제목"],
            ref["저자/출처"],
            ref["URL"],
            ref["발행일"],
            ref["요약"],
            ref["관련 섹션"],
        ]

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.border = THIN_BORDER

            if col_idx == 1:
                cell.font = BOLD_FONT
                cell.alignment = CENTER
            elif col_idx == 4:
                cell.font = URL_FONT
                cell.alignment = LEFT
                cell.hyperlink = value
            else:
                cell.font = BODY_FONT
                cell.alignment = LEFT

        ws.row_dimensions[row_idx].height = 70

    ws.row_dimensions[1].height = 25
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "open-intelligent-ui-references.xlsx")
    wb.save(output_path)
    print(f"[완료] 참고문헌 파일 생성: {output_path}")
    print(f"[정보] 총 {len(references)}개 참고문헌 포함")
    return output_path


if __name__ == "__main__":
    create_references_xlsx()
